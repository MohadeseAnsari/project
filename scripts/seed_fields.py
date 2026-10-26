"""Seed the ``fields`` table from ``data/fields.xlsx``.

Each raw field name is normalized and inserted with
``ON CONFLICT (normalized_name) DO NOTHING`` so re-running the script is
idempotent and duplicate normalized identities are never created.

The field table is only the stable dedup identity for discipline names.
Classification is done by the LLM per record during generation; no category or
keyword is stored on a field.
"""

import asyncio
import sys
from pathlib import Path

import asyncpg
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.config import settings
from app.fields.normalization import normalize_field


async def seed_fields():
    wb = openpyxl.load_workbook("data/fields.xlsx", read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    raw_fields = [str(r[0]) for r in rows[1:] if r and r[0]]

    conn = await asyncpg.connect(settings.database_url.replace("+asyncpg", ""))
    try:
        inserted = 0
        skipped = 0
        for raw in raw_fields:
            norm = normalize_field(raw)

            status = await conn.execute(
                """
                INSERT INTO fields (name, normalized_name, source)
                VALUES ($1, $2, 'spreadsheet')
                ON CONFLICT (normalized_name) DO NOTHING
                """,
                raw,
                norm,
            )
            # asyncpg command tag: "INSERT 0 1" = row written, "INSERT 0 0" = conflict-skipped
            if status == "INSERT 0 1":
                inserted += 1
            else:
                skipped += 1

        print(f"Inserted: {inserted}, Skipped: {skipped}")
    finally:
        await conn.close()


if __name__ == "__main__":
    asyncio.run(seed_fields())
