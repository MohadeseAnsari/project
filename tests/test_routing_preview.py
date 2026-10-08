"""Tests for the auditable routing chain: preview, management, catalogue coverage."""

import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.db.models import Category, CategoryChannelRoute, Channel, Field, Main
from app.main import app as api_app
from app.panel.app import panel_app
from app.panel.auth import hash_password


async def _admin(session: AsyncSession) -> None:
    from app.db.models import PanelUser, Role

    role = (await session.execute(select(Role).where(Role.name == "admin"))).scalar_one_or_none()
    if role is None:
        role = Role(name="admin", description="Full access")
        session.add(role)
        await session.flush()
    session.add(
        PanelUser(username="rout-admin", password_hash=hash_password("secret123"),
                  role_id=role.id, is_active=True)
    )
    await session.commit()


def _api() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=api_app), base_url="http://test")


def _panel() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=panel_app), base_url="http://test")


@pytest.mark.asyncio
async def test_routing_preview_for_field(session: AsyncSession):
    """Raw field name resolves to category, keyword, and live channels."""
    channel = Channel(platform="bale", chat_id=71001, username="@c", is_active=True)
    session.add(channel)
    await session.commit()
    async with _api() as client:
        response = await client.get(
            "/routing/preview",
            params={"shared_Secret": settings.shared_secret, "field": "مدیریت موزه"},
        )
    assert response.status_code == 200
    body = response.json()
    assert body["category"] == "هنر و معماری"
    assert body["matched_keyword"] == "موزه"
    assert len(body["channels"]) == 1
    assert body["channels"][0]["via"] in ("category", "field", "fallback")


@pytest.mark.asyncio
async def test_routing_preview_validation(session: AsyncSession):
    """Auth, exactly-one-input rule, and unknown records are enforced."""
    async with _api() as client:
        assert (await client.get("/routing/preview")).status_code == 401
        secret = {"shared_Secret": settings.shared_secret}
        both = await client.get("/routing/preview", params={**secret, "field": "x", "record_id": 1})
        assert both.status_code == 422
        neither = await client.get("/routing/preview", params=secret)
        assert neither.status_code == 422
        missing = await client.get("/routing/preview", params={**secret, "record_id": 999999})
        assert missing.status_code == 404


@pytest.mark.asyncio
async def test_routing_preview_for_record(session: AsyncSession):
    """A record previews its snapshot category plus live channels."""
    cat = Category(name="پزشکی و علوم سلامت", description="Medical")
    session.add(cat)
    await session.commit()
    field = Field(name="f", normalized_name="f", category_id=cat.id, matched_keyword="پزشکی")
    session.add(field)
    await session.commit()
    channel = Channel(platform="bale", chat_id=71002, username="@c", is_active=True)
    session.add(channel)
    await session.commit()
    session.add(CategoryChannelRoute(category_id=cat.id, channel_id=channel.id))
    await session.commit()
    row = Main(
        username="t", title="T", abstract="A", field="f", university="U",
        author="A", supervisor="S", status="review",
        field_id=field.id, category_id=cat.id,
        output_json={"title": "T", "lead": "L", "body": "B"},
    )
    session.add(row)
    await session.commit()
    async with _api() as client:
        response = await client.get(
            "/routing/preview",
            params={"shared_Secret": settings.shared_secret, "record_id": row.id},
        )
    assert response.status_code == 200
    body = response.json()
    assert body["category"] == "پزشکی و علوم سلامت"
    assert body["matched_keyword"] == "پزشکی"
    assert body["channels"][0]["via"] == "category"


@pytest.mark.asyncio
async def test_panel_routing_page_and_category_management(session: AsyncSession):
    """Panel renders the preview and admins manage category routes."""
    await _admin(session)
    cat = Category(name="پزشکی و علوم سلامت", description="Medical")
    session.add(cat)
    await session.commit()
    channel = Channel(platform="bale", chat_id=71003, username="@c", is_active=True)
    session.add(channel)
    await session.commit()
    async with _panel() as client:
        login = await client.post(
            "/login", data={"username": "rout-admin", "password": "secret123"},
            follow_redirects=False,
        )
        assert login.status_code == 302
        page = await client.get("/routing", params={"field": "رادیولوژی"})
        assert page.status_code == 200
        assert "رادیولوژی" in page.text
        added = await client.post(
            "/channels/category-route",
            data={"category_id": cat.id, "channel_id": channel.id},
            follow_redirects=False,
        )
        assert added.status_code == 302
        preview = await client.get("/routing", params={"field": "رادیولوژی"})
        assert "دسته" in preview.text
        removed = await client.post(
            "/channels/category-unroute",
            data={"category_id": cat.id, "channel_id": channel.id},
            follow_redirects=False,
        )
        assert removed.status_code == 302
    remaining = (
        await session.execute(select(CategoryChannelRoute))
    ).scalars().all()
    assert remaining == []


@pytest.mark.asyncio
async def test_record_detail_shows_snapshot_vs_current(session: AsyncSession):
    """The review page shows both the stored and the current category."""
    await _admin(session)
    old_cat = Category(name="مهندسی و فناوری", description="Old")
    new_cat = Category(name="پزشکی و علوم سلامت", description="New")
    session.add_all([old_cat, new_cat])
    await session.commit()
    field = Field(name="f", normalized_name="f", category_id=new_cat.id, matched_keyword="بهداشت")
    session.add(field)
    await session.commit()
    row = Main(
        username="t", title="T", abstract="A", field="f", university="U",
        author="A", supervisor="S", status="review",
        field_id=field.id, category_id=old_cat.id,
        output_json={"title": "T", "lead": "L", "body": "B"},
    )
    session.add(row)
    await session.commit()
    async with _panel() as client:
        await client.post(
            "/login", data={"username": "rout-admin", "password": "secret123"},
            follow_redirects=False,
        )
        page = await client.get(f"/records/{row.id}")
        assert page.status_code == 200
        assert "دستهٔ ثبت‌شده" in page.text
        assert "دستهٔ کنونی رشته" in page.text
        assert "تغییر کرده است" in page.text


def _catalogue_names() -> list[str]:
    """All spreadsheet catalogue names (repo data, no external fixture)."""
    import openpyxl

    wb = openpyxl.load_workbook("data/fields.xlsx", read_only=True)
    ws = wb[wb.sheetnames[0]]
    return [str(r[0]) for r in list(ws.iter_rows(values_only=True))[1:] if r and r[0]]


def test_catalogue_yields_zero_unexpected_other():
    """Every catalogue name lands somewhere sensible; only reviewed degree labels stay other."""
    from app.fields.classification import classify_field
    from app.fields.normalization import normalize_field

    reviewed_other = {"دکترای تخصصی پژوهشی", "کارشناسی حرفه‌ای ارشد"}
    others = [
        name for name in _catalogue_names()
        if classify_field(normalize_field(name)) == "other"
    ]
    assert set(others) == reviewed_other, f"unexpected other: {sorted(set(others) - reviewed_other)}"


def test_catalogue_classification_is_deterministic():
    """Two full passes over the catalogue produce identical categories."""
    from app.fields.classification import classify_field
    from app.fields.normalization import normalize_field

    names = _catalogue_names()
    first = [classify_field(normalize_field(n)) for n in names]
    second = [classify_field(normalize_field(n)) for n in names]
    assert first == second


@pytest.mark.asyncio
async def test_seed_twice_produces_identical_categories(session: AsyncSession):
    """Re-running the seed never changes an existing field's category."""
    import importlib.util
    from pathlib import Path

    from sqlalchemy import func

    # The per-test truncate wipes the migration-seeded categories, so plant
    # the 14 known names first (exactly as 002_seed_categories.sql does).
    for name in (
        "پزشکی و علوم سلامت", "مهندسی و فناوری", "هنر و معماری", "دامپزشکی",
        "علوم اجتماعی و اقتصاد", "علوم انسانی", "علوم پایه",
        "علوم تربیتی و روان‌شناسی", "کامپیوتر و فناوری اطلاعات",
        "حقوق و علوم سیاسی", "کشاورزی و منابع طبیعی",
        "ورزش و تربیت بدنی", "رسانه و ارتباطات", "other",
    ):
        session.add(Category(name=name, description=name))
    await session.commit()

    spec = importlib.util.spec_from_file_location(
        "seed_fields", str(Path("scripts/seed_fields.py").resolve())
    )
    seed_module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(seed_module)

    await seed_module.seed_fields()
    first = dict(
        (await session.execute(select(Field.normalized_name, Category.name).join(Category, Category.id == Field.category_id))).all()
    )
    assert len(first) > 500
    await seed_module.seed_fields()
    second = dict(
        (await session.execute(select(Field.normalized_name, Category.name).join(Category, Category.id == Field.category_id))).all()
    )
    assert first == second
    keyworded = (
        await session.execute(select(func.count()).select_from(Field).where(Field.matched_keyword.isnot(None)))
    ).scalar_one()
    # Every classified field stores its reason; only reviewed `other` rows
    # (degree labels, not fields of study) legitimately have none.
    others = sorted(name for name, cat in second.items() if cat == "other")
    assert keyworded == len(second) - len(others)
    assert others == ["دکترای تخصصی پژوهشی", "کارشناسی حرفه‌ای ارشد"]
